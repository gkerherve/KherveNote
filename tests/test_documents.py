import pytest

from khervenote import documents


@pytest.fixture
def pdf(tmp_path):
    import pymupdf
    doc = pymupdf.open()
    texts = ["Introduction\nXPS measures binding energies.",
             "Backgrounds\nThe Shirley background is iterative.\nTougaard uses a loss function.",
             "Peak shapes\nVoigt profiles mix Gaussian and\nLorentzian broadening."]
    for t in texts:
        page = doc.new_page()
        page.insert_text((72, 72), t)
    doc.set_toc([[1, "Introduction", 1], [1, "Backgrounds", 2], [2, "Shirley", 2],
                 [1, "Peak shapes", 3]])
    path = tmp_path / "xps.pdf"
    doc.save(path)
    return path


def test_pdf_sections_follow_the_bookmarks(pdf):
    doc = documents.read(pdf)
    assert [(s.title, s.level, s.page) for s in doc.sections] == [
        ("Introduction", 1, 1), ("Backgrounds", 1, 2), ("Shirley", 2, 2), ("Peak shapes", 1, 3)]
    # Two bookmarks on page 2 split it at the second title: no text twice.
    assert sum("iterative" in sec.text for sec in doc.sections) == 1
    assert doc.sections[2].text.startswith("Shirley")
    assert "Voigt" in doc.sections[3].text
    assert doc.sections[0].text.startswith("Introduction")


def test_pdf_without_bookmarks_has_pages(tmp_path):
    import pymupdf
    d = pymupdf.open()
    d.new_page().insert_text((72, 72), "one")
    d.new_page().insert_text((72, 72), "two")
    d.save(tmp_path / "p.pdf")
    doc = documents.read(tmp_path / "p.pdf")
    assert [s.title for s in doc.sections] == ["Page 1", "Page 2"]


def test_find_matches_across_line_breaks(pdf):
    doc = documents.read(pdf)
    hits = documents.find(doc, "gaussian and lorentzian")
    assert len(hits) == 1 and hits[0].where == "p. 3"
    h = hits[0]
    assert h.snippet[h.start:h.start + h.length].lower() == "gaussian and lorentzian"
    assert documents.find(doc, "   ") == []


def test_word_headings_paragraphs_and_tables(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Lead paragraph.")
    d.add_heading("Methods", level=1)
    d.add_paragraph("We used Al Ka.")
    d.add_heading("Details", level=2)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Pass energy", "20 eV"
    d.save(tmp_path / "n.docx")
    doc = documents.read(tmp_path / "n.docx")
    assert [(s.title, s.level) for s in doc.sections] == [("", 1), ("Methods", 1), ("Details", 2)]
    assert doc.sections[2].text == "Pass energy | 20 eV"
    assert doc.kind == "word"


def test_slides_one_section_each_with_notes(tmp_path):
    import pptx
    deck = pptx.Presentation()
    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = "Why XPS"
    s.placeholders[1].text_frame.text = "Surface sensitive"
    s.notes_slide.notes_text_frame.text = "Mention 10 nm"
    deck.save(tmp_path / "talk.pptx")
    doc = documents.read(tmp_path / "talk.pptx")
    assert doc.sections[0].title == "Why XPS" and doc.sections[0].page == 1
    assert "- Surface sensitive" in doc.sections[0].text
    assert "Speaker notes: Mention 10 nm" in doc.sections[0].text


def test_markdown_and_unknown(tmp_path):
    (tmp_path / "a.md").write_text("intro\n# One\nbody\n## Two\nmore")
    doc = documents.read(tmp_path / "a.md")
    assert [(s.title, s.level, s.text) for s in doc.sections] == [
        ("", 1, "intro"), ("One", 1, "body"), ("Two", 2, "more")]
    with pytest.raises(ValueError):
        documents.read(tmp_path / "old.doc")


def test_best_passages_pick_the_relevant_part():
    doc = documents.Document("big", "text", [
        documents.DocSection(f"Part {i}", ("filler words about nothing. " * 60)) for i in range(30)])
    doc.sections[17].text += " The Shirley background subtracts inelastic electrons."
    chosen = documents.best_passages(doc, "How does the Shirley background work?", budget=3000)
    assert any("Shirley" in p.text for p in chosen)
    assert sum(len(p.text) for p in chosen) <= 3000
    assert len(documents.batches(doc, budget=5000)) > 5
