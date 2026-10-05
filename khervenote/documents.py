# KherveNote — attached documents
# Copyright (C) 2026  Gwilherm Kerherve
# SPDX-License-Identifier: GPL-3.0-or-later
"""Read PDF, Word, PowerPoint and text documents attached to a note into
sections of plain text, search them, and pick the passages that matter
for a question.  Qt-free.

Sections come from the document's own structure where it has one — a
PDF's bookmarks, Word's heading styles, one per slide — and otherwise
one per page, so a hit can always say where it is.
"""
from __future__ import annotations

import math
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

#: Suffix → kind, for the icon and the reader.
KINDS = {".pdf": "pdf", ".docx": "word", ".pptx": "slides", ".txt": "text",
         ".md": "text", ".markdown": "text"}
FILTER = "Documents (*.pdf *.docx *.pptx *.txt *.md)"


@dataclass
class DocSection:
    title: str
    text: str
    level: int = 1
    page: Optional[int] = None        # 1-based, where known

    @property
    def where(self) -> str:
        return f"p. {self.page}" if self.page else self.title


@dataclass
class Document:
    name: str
    kind: str
    sections: list[DocSection] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(f"{s.title}\n{s.text}" if s.title else s.text
                           for s in self.sections)


def kind_of(path: Path) -> Optional[str]:
    return KINDS.get(Path(path).suffix.lower())


def read(path: Path, name: str = "") -> Document:
    path = Path(path)
    kind = kind_of(path)
    if kind is None:
        raise ValueError(f"{path.suffix or 'This'} files cannot be read — use PDF, "
                         "Word (.docx), PowerPoint (.pptx) or text.")
    doc = Document(name or path.name, kind)
    doc.sections = {"pdf": _read_pdf, "word": _read_docx, "slides": _read_pptx,
                    "text": _read_text}[kind](path)
    doc.sections = [s for s in doc.sections if s.text.strip() or s.title.strip()]
    return doc


def _clean(text: str) -> str:
    text = text.replace("­", "")
    text = re.sub(r"-\n(?=[a-z])", "", text)           # hyphenated line breaks
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _read_pdf(path: Path) -> list[DocSection]:
    import pymupdf
    with pymupdf.open(path) as pdf:
        pages = [_clean(p.get_text("text")) for p in pdf]
        toc = pdf.get_toc(simple=True)
    if not toc:
        return [DocSection(f"Page {i + 1}", t, 1, i + 1) for i, t in enumerate(pages)]
    marks = [(lvl, title.strip(), page) for lvl, title, page in toc
             if 0 < page <= len(pages)]
    # One string for the whole document, so a section can start where its
    # title is printed rather than at the top of its page — several
    # bookmarks on one page then split it instead of each getting all of it.
    full, starts = "", []
    for t in pages:
        starts.append(len(full))
        full += t + "\n\n"
    flat = full.lower()
    cuts = []
    for lvl, title, page in marks:
        lo = max(starts[page - 1], cuts[-1] if cuts else 0)
        hi = starts[page] if page < len(pages) else len(full)
        at = _find_title(flat, title.lower(), lo, hi)
        cuts.append(at if at >= 0 else lo)
    out = []
    if cuts and cuts[0] > 0 and full[:cuts[0]].strip():
        out.append(DocSection("Before the first section", full[:cuts[0]].strip(), 1, 1))
    for i, (lvl, title, page) in enumerate(marks):
        end = cuts[i + 1] if i + 1 < len(cuts) else len(full)
        out.append(DocSection(title, full[cuts[i]:end].strip(), lvl, page))
    if not _useful(out):
        # Bookmarks that all carry the paper's title, or mostly cover
        # nothing, are worse than none: use the pages.
        return [DocSection(f"Page {i + 1}", t, 1, i + 1) for i, t in enumerate(pages)]
    return out


def _useful(sections: list[DocSection]) -> bool:
    if not sections:
        return False
    titles = {s.title.strip().lower() for s in sections}
    empty = sum(1 for s in sections if len(s.text) < 40)
    return len(titles) >= 0.6 * len(sections) and empty <= 0.4 * len(sections)


def _find_title(flat: str, title: str, lo: int, hi: int) -> int:
    """Where *title* is printed between *lo* and *hi*, allowing for line
    breaks and spacing that differ from the bookmark."""
    words = title.split()
    if not words:
        return -1
    m = re.compile(r"\s+".join(re.escape(w) for w in words)).search(flat, lo, hi)
    return m.start() if m else -1


def _read_docx(path: Path) -> list[DocSection]:
    import docx
    d = docx.Document(str(path))
    out = [DocSection("", "", 1)]
    for block in _docx_blocks(d):
        if isinstance(block, tuple):
            title, level = block
            out.append(DocSection(title, "", level))
        elif block.strip():
            out[-1].text += ("\n" if out[-1].text else "") + block
    return out


def _docx_blocks(d):
    """Paragraphs and tables in document order; headings as (title, level)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            if m and p.text.strip():
                yield (p.text.strip(), int(m.group(1)))
            elif style == "Title" and p.text.strip():
                yield (p.text.strip(), 1)
            else:
                yield p.text
        elif tag == "tbl":
            t = Table(child, d)
            for row in t.rows:
                yield " | ".join(c.text.strip() for c in row.cells)


def _read_pptx(path: Path) -> list[DocSection]:
    import pptx
    deck = pptx.Presentation(str(path))
    out = []
    for i, slide in enumerate(deck.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        lines = []
        for shape in slide.shapes:
            if shape == slide.shapes.title or not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs).strip()
                if text:
                    lines.append("  " * para.level + "- " + text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("Speaker notes: " + notes)
        out.append(DocSection(title or f"Slide {i}", "\n".join(lines), 1, i))
    return out


def _read_text(path: Path) -> list[DocSection]:
    out = [DocSection("", "", 1)]
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            out.append(DocSection(m.group(2).strip(), "", len(m.group(1))))
        else:
            out[-1].text += line + "\n"
    for s in out:
        s.text = s.text.strip()
    return out


# ── search ──────────────────────────────────────────────────────────

@dataclass
class Hit:
    section: int
    where: str
    snippet: str          # the sentence(s) around the match
    start: int            # match position inside the snippet
    length: int


def find(doc: Document, query: str, context: int = 140, limit: int = 200) -> list[Hit]:
    """Every place *query* occurs (case-insensitive; spaces match any
    whitespace, so a sentence split over lines is still found)."""
    words = query.split()
    if not words:
        return []
    pattern = re.compile(r"\s+".join(re.escape(w) for w in words), re.IGNORECASE)
    hits = []
    for i, sec in enumerate(doc.sections):
        text = sec.text
        for m in pattern.finditer(text):
            a = max(0, m.start() - context)
            b = min(len(text), m.end() + context)
            # Widen to whole words.
            while a > 0 and not text[a - 1].isspace():
                a -= 1
            while b < len(text) and not text[b].isspace():
                b += 1
            raw = text[a:b]
            snippet = re.sub(r"\s+", " ", raw).strip()
            inner = re.sub(r"\s+", " ", text[a:m.start()]).lstrip()
            where = sec.where if sec.page or sec.title else doc.name
            hits.append(Hit(i, where,
                            ("… " if a > 0 else "") + snippet + (" …" if b < len(text) else ""),
                            len(inner) + (2 if a > 0 else 0),
                            len(re.sub(r"\s+", " ", m.group(0)))))
            if len(hits) >= limit:
                return hits
    return hits


# ── passages for the AI ─────────────────────────────────────────────

_WORD = re.compile(r"\w+", re.UNICODE)
_STOP = set("""a an and are as at be by for from has have how i in is it its of on
or that the this to was were what when where which who why will with do does
did can could should would about into than then there these those not no you
your we our they their le la les un une des et est du de en que qui dans pour
""".split())


def _terms(text: str) -> list[str]:
    return [w for w in (m.group(0).lower() for m in _WORD.finditer(text))
            if w not in _STOP and len(w) > 1]


@dataclass
class Passage:
    where: str
    text: str


def passages(doc: Document, size: int = 1400) -> list[Passage]:
    """The document cut into passages of about *size* characters, never
    across sections, each knowing where it comes from."""
    out = []
    for sec in doc.sections:
        paras = [p for p in re.split(r"\n\s*\n|\n(?=- )", sec.text) if p.strip()]
        buf = ""
        for p in paras:
            if buf and len(buf) + len(p) > size:
                out.append(Passage(sec.where, buf.strip()))
                buf = ""
            buf += p + "\n"
            while len(buf) > size * 1.5:
                out.append(Passage(sec.where, buf[:size].strip()))
                buf = buf[size:]
        if buf.strip() or (sec.title and not out):
            out.append(Passage(sec.where, ((sec.title + "\n") if sec.title else "") + buf.strip()))
    return out


def best_passages(doc: Document, question: str, budget: int = 9000) -> list[Passage]:
    """The passages most relevant to *question* (BM25), in document
    order, within *budget* characters — what a small local model can
    read in one go."""
    chunks = passages(doc)
    if sum(len(c.text) for c in chunks) <= budget:
        return chunks
    q = _terms(question)
    docs = [_terms(c.text) for c in chunks]
    n = len(docs)
    avg = sum(len(d) for d in docs) / max(1, n)
    df = Counter(t for d in docs for t in set(d))
    scores = []
    for i, d in enumerate(docs):
        tf = Counter(d)
        s = 0.0
        for t in q:
            if t not in tf:
                continue
            idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
            s += idf * tf[t] * 2.2 / (tf[t] + 1.2 * (0.25 + 0.75 * len(d) / max(1, avg)))
        scores.append((s, i))
    chosen, used = [], 0
    for s, i in sorted(scores, reverse=True):
        if s <= 0 and chosen:
            break
        if used + len(chunks[i].text) > budget:
            continue
        chosen.append(i)
        used += len(chunks[i].text)
    return [chunks[i] for i in sorted(chosen)]


def batches(doc: Document, budget: int = 9000) -> list[str]:
    """The whole document in pieces of at most *budget* characters, for
    summarising part by part."""
    out, buf = [], ""
    for p in passages(doc):
        piece = f"[{p.where}]\n{p.text}\n\n"
        if buf and len(buf) + len(piece) > budget:
            out.append(buf)
            buf = ""
        buf += piece
    if buf:
        out.append(buf)
    return out
